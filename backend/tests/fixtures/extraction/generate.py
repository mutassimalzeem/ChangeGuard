"""Regenerate synthetic fixture files; requires the backend dependencies only."""

from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

ROOT = Path(__file__).parent


def pdf(pages):
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_id = writer._add_object(font)
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_id})}
        )
        stream = DecodedStreamObject()
        operators = ["BT /F1 12 Tf 50 740 Td"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            operators += [f"({escaped}) Tj", "0 -20 Td"]
        operators += ["ET"]
        stream.set_data("\n".join(operators).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_metadata({"/Author": "Fixture Engineer"})
    return writer


pdf(
    [["1 Scope", "Supply concrete for foundations."], ["2 Payment", "Payment within thirty days."]]
).write(ROOT / "contract.pdf")
image = Image.new("RGB", (1200, 450), "white")
draw = ImageDraw.Draw(image)
font = ImageFont.load_default(size=48)
draw.text((40, 50), "Daily Report", fill="black", font=font)
draw.text((40, 150), "Weather: Clear", fill="black", font=font)
draw.text((40, 250), "Activities: Concrete works", fill="black", font=font)
image.save(ROOT / "report.png")
image.save(ROOT / "scanned.pdf", "PDF", resolution=150)
mixed = pdf([["1 Scope", "Native text must survive mixed-page OCR."]])
mixed.add_page(PdfReader(ROOT / "scanned.pdf").pages[0])
mixed.write(ROOT / "mixed.pdf")
hybrid = PdfWriter()
scan_page = PdfReader(ROOT / "scanned.pdf").pages[0]
overlay = pdf([["Native header"]]).pages[0]
overlay.scale_to(float(scan_page.mediabox.width), float(scan_page.mediabox.height))
scan_page.merge_page(overlay)
hybrid.add_page(scan_page)
hybrid.write(ROOT / "hybrid.pdf")
doc = Document()
doc.core_properties.author = "Fixture Author"
doc.core_properties.created = datetime(2026, 1, 15)
doc.core_properties.revision = 3
doc.add_heading("1 Scope", level=1)
doc.add_paragraph("Supply concrete and reinforcement. " * 24)
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text, table.cell(0, 1).text = "Topic", "Value"
table.cell(1, 0).text, table.cell(1, 1).text = "Grade", "C30"
doc.add_heading("2 Payment", level=1)
doc.add_paragraph("Payment within thirty days.")
doc.save(ROOT / "contract.docx")
book = Workbook()
book.properties.creator = "Fixture QS"
book.properties.created = datetime(2026, 1, 16)
book.properties.revision = "2"
sheet = book.active
sheet.title = "BOQ"
sheet.append(["Item", "Description", "Quantity", "Unit", "Rate", "Amount"])
sheet.append(["1.01", "Concrete foundations\nIncludes placing and curing", 10, "m3", 100, "=C2*E2"])
sheet.append(["1.02", "Reinforcement", 2, "t", 500, 1000])
book.save(ROOT / "boq.xlsx")
