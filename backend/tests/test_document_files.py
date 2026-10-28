"""Validate actual supported formats and immutable storage writes."""

from io import BytesIO
from unittest.mock import MagicMock
from zipfile import ZipFile

import pytest
from botocore.exceptions import ClientError
from PIL import Image

from app.domain.services.document_files import TYPES, validate_document
from app.domain.services.object_storage import ObjectStorageService, StorageError


@pytest.mark.parametrize("extension", ["png", "jpg", "jpeg"])
def test_valid_images(extension):
    output = BytesIO()
    Image.new("RGB", (2, 2)).save(output, format="PNG" if extension == "png" else "JPEG")
    assert validate_document("image." + extension, output.getvalue())[0] == extension


@pytest.mark.parametrize("extension", ["docx", "xlsx", "zip"])
def test_valid_archives(extension):
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(
            "word/document.xml" if extension == "docx" else "xl/workbook.xml", "<document/>"
        )
    assert validate_document("document." + extension, output.getvalue())[0] == extension


def test_valid_email():
    assert (
        validate_document("email.eml", b"From: test@example.com\r\nSubject: Change\r\n\r\nBody")[0]
        == "eml"
    )


@pytest.mark.parametrize("extension", TYPES)
def test_corrupt_content_rejected(extension):
    with pytest.raises(ValueError):
        validate_document("bad." + extension, b"invalid document content")


def test_archive_traversal_rejected():
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("../escape", b"file")
    with pytest.raises(ValueError):
        validate_document("unsafe.zip", output.getvalue())


def test_original_is_write_once():
    storage = ObjectStorageService()
    storage._client = MagicMock()
    stream = BytesIO(b"original")
    storage.store_original("a/key", stream, "application/pdf")
    storage._client.put_object.assert_called_once_with(
        Bucket=storage._bucket,
        Key="a/key",
        Body=stream,
        ContentType="application/pdf",
        IfNoneMatch="*",
    )
    storage._client.put_object.side_effect = ClientError(
        {"Error": {"Code": "PreconditionFailed"}}, "PutObject"
    )
    with pytest.raises(StorageError):
        storage.store_original("a/key", stream, "application/pdf")


def test_real_office_documents():
    from docx import Document
    from openpyxl import Workbook

    word = BytesIO()
    document = Document()
    document.add_paragraph("Change order")
    document.save(word)
    assert validate_document("document.docx", word.getvalue())[0] == "docx"
    excel = BytesIO()
    Workbook().save(excel)
    assert validate_document("schedule.xlsx", excel.getvalue())[0] == "xlsx"
